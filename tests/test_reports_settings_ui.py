from dataclasses import replace
from pathlib import Path
import time

from openpyxl import load_workbook
from pypdf import PdfReader
import pytest
from PySide6.QtCore import QEventLoop, QTimer
from PySide6.QtWidgets import QFileDialog, QMessageBox

from smartpower.application import create_application
from smartpower.data.demo import generate_demo
from smartpower.settings.manager import SettingsManager
from smartpower.storage.repository import DatasetRepository
from smartpower.ui.pages.reports import ReportsPage
from smartpower.ui.pages.settings import SettingsPage


def wait_idle(page):
    if page.busy:
        loop = QEventLoop()
        timeout = QTimer()
        timeout.setSingleShot(True)
        timeout.timeout.connect(loop.quit)
        page.operation_finished.connect(loop.quit)
        timeout.start(30000)
        loop.exec()
        timeout.stop()
        page.operation_finished.disconnect(loop.quit)
    assert not page.busy


@pytest.fixture
def page(tmp_path):
    app = create_application([])
    widget = ReportsPage(repository=DatasetRepository(tmp_path / "报告.sqlite3"),
                         manager=SettingsManager(tmp_path / "设置.json", tmp_path))
    widget.resize(850, 850)
    widget.show()
    yield widget
    wait_idle(widget)
    widget.close()
    widget.deleteLater()
    app.processEvents()


def select_saved(page, count=10):
    identifier = page.repository.save(generate_demo(count), "报告模拟数据")
    page.refresh_datasets()
    wait_idle(page)
    page.dataset_combo.setCurrentIndex(page.dataset_combo.findData(identifier))
    return identifier


def test_no_dataset_selection_is_safe(page, monkeypatch):
    def never(*args, **kwargs):
        pytest.fail("未选择数据时不能弹出路径对话框")
    monkeypatch.setattr(QFileDialog, "getSaveFileName", never)
    page.pdf_button.click()
    assert "请先选择" in page.message.text() and not page.busy


def test_stale_font_setting_does_not_silently_switch_to_reference(page, tmp_path, monkeypatch):
    select_saved(page)
    font = tmp_path / "原先配置.ttf"
    font.write_bytes(b"test-only-font-file")
    page.manager.save(replace(page.manager.defaults, font_path=str(font)))
    font.unlink()
    def never(*args, **kwargs):
        pytest.fail("失效字体配置应在弹出路径对话框前提示修复")
    monkeypatch.setattr(QFileDialog, "getSaveFileName", never)
    page.start_export("pdf")
    assert "字体文件不存在" in page.message.text() and "系统设置" in page.message.text()
    assert not page.busy


@pytest.mark.parametrize("kind", ["pdf", "xlsx"])
def test_real_report_export_and_default_path(page, tmp_path, monkeypatch, kind):
    identifier = select_saved(page)
    path = tmp_path / f"中文报告.{kind}"
    dialogs = []
    def choose(*args, **kwargs):
        dialogs.append(args)
        return str(path), ""
    monkeypatch.setattr(QFileDialog, "getSaveFileName", choose)
    page.start_export(kind)
    assert page.busy and not page.pdf_button.isEnabled()
    wait_idle(page)
    assert page.current_snapshot.analysis.dataset_id == identifier
    assert str(path) in page.message.text() and "报告已保存" in page.message.text()
    assert Path(dialogs[0][2]).parent == tmp_path
    assert page.current_snapshot.detection is None
    if kind == "pdf":
        assert "报告模拟数据" in "\n".join(p.extract_text() for p in PdfReader(path).pages)
    else:
        book = load_workbook(path)
        assert book["原始数据"].max_row == 11
        book.close()


@pytest.mark.parametrize("kind", ["pdf", "xlsx"])
def test_export_error_clears_state_and_preserves_file(page, tmp_path, monkeypatch, kind):
    select_saved(page)
    path = tmp_path / f"已有.{kind}"
    path.write_bytes(b"original")
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **kw: (str(path), ""))
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **kw: QMessageBox.StandardButton.Yes)
    def fail(*args, **kwargs):
        raise OSError("测试磁盘错误")
    monkeypatch.setattr("smartpower.reporting.files.os.replace", fail)
    page.start_export(kind)
    wait_idle(page)
    assert "失败" in page.message.text() and page.current_snapshot is None
    assert path.read_bytes() == b"original" and page.pdf_button.isEnabled()


def test_overwrite_requires_confirmation_and_default_no(page, tmp_path, monkeypatch):
    select_saved(page)
    path = tmp_path / "已有.xlsx"
    path.write_bytes(b"original")
    calls = []
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **kw: (str(path), ""))
    def deny(*args):
        calls.append(args)
        return QMessageBox.StandardButton.No
    monkeypatch.setattr(QMessageBox, "question", deny)
    page.start_export("xlsx")
    assert len(calls) == 1 and calls[0][-1] == QMessageBox.StandardButton.No
    assert path.read_bytes() == b"original" and not page.busy


def test_switch_while_dialog_open_cancels(page, tmp_path, monkeypatch):
    select_saved(page)
    path = tmp_path / "取消.pdf"
    def change(*args, **kwargs):
        page.report_title.setText("对话框中改标题")
        return str(path), ""
    monkeypatch.setattr(QFileDialog, "getSaveFileName", change)
    page.start_export("pdf")
    assert not path.exists() and "已取消" in page.message.text()


def test_worker_uses_original_options_and_keeps_event_loop_alive(page, tmp_path, monkeypatch):
    from smartpower.reporting.service import build_report_snapshot
    identifier = select_saved(page)
    path = tmp_path / "快照.xlsx"
    page.report_title.setText("启动时标题")
    page.detection_enabled.setChecked(True)
    for edit, value in zip(page.threshold_inputs, ("210", "230", "20")):
        edit.setText(value)
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **kw: (str(path), ""))
    def delayed(*args, **kwargs):
        time.sleep(0.1)
        return build_report_snapshot(*args, **kwargs)
    monkeypatch.setattr("smartpower.ui.pages.reports.build_report_snapshot", delayed)
    ticks = []
    timer = QTimer()
    timer.setInterval(5)
    timer.timeout.connect(lambda: ticks.append(1))
    timer.start()
    page.start_export("xlsx")
    page.dataset_combo.setCurrentIndex(-1)  # Programmatic race despite disabled controls.
    page.report_title.setText("后来的标题")
    page.voltage_upper.setText("999")
    wait_idle(page)
    timer.stop()
    assert len(ticks) >= 2 and page.current_snapshot is None
    assert "启动时" in page.message.text() and "报告模拟数据" in page.message.text()
    assert "尚未选择" in page.metadata.toPlainText()
    book = load_workbook(path)
    overview = {r[0]: r[1] for r in book["报告概况"].iter_rows(values_only=True)}
    anomaly = {r[0]: r[1] for r in list(book["异常检测"].iter_rows(values_only=True))[:12]}
    assert overview["报告标题"] == "启动时标题" and overview["数据集 ID"] == identifier
    assert anomaly["电压上限（V）"] == 230
    book.close()


@pytest.mark.parametrize("value", ["", "NaN", "Infinity", "-1", "240"])
def test_invalid_thresholds_fail_before_export(page, value):
    select_saved(page)
    page.detection_enabled.setChecked(True)
    for edit, val in zip(page.threshold_inputs, (value, "230", "20")):
        edit.setText(val)
    page.start_export("pdf")
    assert "阈值无效" in page.message.text() and not page.busy


def test_settings_save_reload_restore_confirmed(tmp_path, monkeypatch):
    app = create_application([])
    manager = SettingsManager(tmp_path / "用户设置.json", tmp_path)
    widget = SettingsPage(manager=manager)
    widget.show()
    widget.report_title.setText("用户专业报告")
    widget.startup_page.setCurrentIndex(widget.startup_page.findData("reports"))
    widget.save_button.click()
    assert "已保存" in widget.message.text()
    second = SettingsPage(manager=SettingsManager(manager.path, tmp_path))
    assert second.report_title.text() == "用户专业报告"
    assert second.startup_page.currentData() == "reports"
    monkeypatch.setattr(QMessageBox, "question", lambda *a: QMessageBox.StandardButton.No)
    widget.restore_button.click()
    assert manager.load().value.report_title == "用户专业报告"
    monkeypatch.setattr(QMessageBox, "question", lambda *a: QMessageBox.StandardButton.Yes)
    widget.restore_button.click()
    assert manager.load().value == manager.defaults and "已恢复" in widget.message.text()
    widget.report_directory.setText(str(tmp_path / "不存在"))
    widget.save_button.click()
    assert "不存在" in widget.message.text() and not (tmp_path / "不存在").exists()
    widget.close()
    second.close()
    widget.deleteLater()
    second.deleteLater()
    app.processEvents()


def test_startup_preference_and_live_report_settings(tmp_path, monkeypatch):
    from smartpower.ui.main_window import MainWindow
    from smartpower.ui.paths import default_settings_manager
    monkeypatch.setenv("SMARTPOWER_SETTINGS_PATH", str(tmp_path / "偏好.json"))
    app = create_application([])
    manager = default_settings_manager()
    manager.save(replace(manager.defaults, startup_page="reports", report_title="已保存标题", chart_directory=str(tmp_path)))
    window = MainWindow()
    window.show()
    assert window.stack.currentWidget() is window.pages["reports"]
    assert window.pages["reports"].report_title.text() == "已保存标题"
    settings = window.pages["settings"]
    settings.report_title.setText("实时新标题")
    settings.save_button.click()
    assert window.pages["reports"].report_title.text() == "实时新标题"
    window.close()
    window.deleteLater()
    app.processEvents()


def test_chart_export_uses_saved_default_directory(tmp_path, monkeypatch):
    from smartpower.ui.paths import default_settings_manager
    from smartpower.visualization.plot_widget import TimePlotWidget
    from smartpower.analysis.service import analyze_dataset
    app = create_application([])
    manager = default_settings_manager()
    manager.save(replace(manager.defaults, chart_directory=str(tmp_path)))
    widget = TimePlotWidget("voltage")
    widget.set_result(analyze_dataset(generate_demo(2), "voltage"))
    calls = []
    def cancel(*args, **kwargs):
        calls.append(args)
        return "", ""
    monkeypatch.setattr(QFileDialog, "getSaveFileName", cancel)
    widget.export_png()
    assert Path(calls[0][2]).parent == tmp_path
    widget.close()
    widget.deleteLater()
    app.processEvents()


def test_window_cannot_close_during_report_export(tmp_path, monkeypatch):
    from smartpower.ui.main_window import MainWindow
    from smartpower.reporting.service import build_report_snapshot
    app = create_application([])
    window = MainWindow()
    window.show()
    page = window.pages["reports"]
    page.repository = DatasetRepository(tmp_path / "关闭保护.sqlite3")
    select_saved(page)
    def slow(*args, **kwargs):
        time.sleep(0.1)
        return build_report_snapshot(*args, **kwargs)
    monkeypatch.setattr("smartpower.ui.pages.reports.build_report_snapshot", slow)
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **kw: (str(tmp_path / "报告.xlsx"), ""))
    page.start_export("xlsx")
    window.navigate("home")
    window.close()
    assert window.isVisible() and window.stack.currentWidget() is page
    wait_idle(page)
    assert page.current_snapshot is not None
    window.close()
    assert not window.isVisible()
    window.deleteLater()
    app.processEvents()
