from dataclasses import FrozenInstanceError
import sqlite3

import numpy as np
from openpyxl import load_workbook
from pypdf import PdfReader
import pytest

from smartpower.analysis.service import analyze_saved_dataset
from smartpower.anomaly.models import Thresholds
from smartpower.anomaly.service import detect_saved_dataset
from smartpower.config import APP_NAME
from smartpower.data.demo import generate_demo
from smartpower.data.validation import DataError
from smartpower.reporting.excel_exporter import export_excel
from smartpower.reporting.models import NOT_DETECTED, NO_EVENTS
from smartpower.reporting.pdf_exporter import export_pdf
from smartpower.reporting.service import build_report_snapshot
from smartpower.storage.repository import DatasetRepository


@pytest.fixture
def report(tmp_path):
    repo = DatasetRepository(tmp_path / "报告.sqlite3")
    dataset = generate_demo(5)
    dataset.frame["电压"] = [205., 210., 220., 230., 235.]
    dataset.frame["电流"] = [21., 10., 20., 30., 10.]
    identifier = repo.save(dataset, "中文模拟数据")
    return repo, identifier, build_report_snapshot(repo, identifier, thresholds=Thresholds(210, 230, 20))


def test_snapshot_matches_v03_v04_and_source(report):
    repo, identifier, snapshot = report
    reference = analyze_saved_dataset(repo, identifier)
    detection = detect_saved_dataset(repo, identifier, snapshot.thresholds)
    assert snapshot.analysis.voltage == reference.voltage
    assert snapshot.analysis.current == reference.current
    assert snapshot.analysis.power.active_statistics == reference.power.active_statistics
    assert snapshot.analysis.power.power_factor_statistics == reference.power.power_factor_statistics
    np.testing.assert_array_equal(snapshot.analysis.power.active, reference.power.active)
    assert snapshot.detection == detection
    assert snapshot.source_label == "模拟数据，非真实设备测量"
    assert snapshot.product_name == APP_NAME and snapshot.version == "1.0.0"


def test_one_read_immutable_snapshot_and_no_database_changes(report, monkeypatch):
    repo, identifier, first = report
    with sqlite3.connect(repo.path) as connection:
        before = tuple(connection.iterdump())
    load, calls = repo.load, []
    def once(identifier):
        calls.append(identifier)
        return load(identifier)
    monkeypatch.setattr(repo, "load", once)
    snapshot = build_report_snapshot(repo, identifier)
    assert calls == [identifier]
    assert snapshot.report_id != first.report_id
    assert snapshot.detection is None and snapshot.detection_status == NOT_DETECTED
    with pytest.raises(FrozenInstanceError):
        snapshot.title = "修改"
    for array in (snapshot.analysis.times, snapshot.analysis.samples.voltage, snapshot.power_factors,
                  snapshot.analysis.power.active):
        with pytest.raises(ValueError):
            array.setflags(write=True)
    with sqlite3.connect(repo.path) as connection:
        assert tuple(connection.iterdump()) == before


def test_switch_datasets_keeps_old_snapshot_consistent(report):
    repo, first, old = report
    dataset = generate_demo(2)
    dataset.simulated = False
    dataset.frame["电压"] = 110.
    second = repo.save(dataset, "第二份来源未核验")
    new = build_report_snapshot(repo, second)
    assert new.analysis.dataset_id == second and old.analysis.dataset_id == first
    assert new.analysis.voltage.mean == 110 and old.analysis.count == 5
    assert "来源未核验" in new.source_label and "模拟数据" in old.source_label
    assert old.detection.dataset_id == first


def test_detected_zero_differs_from_not_executed(report):
    repo, identifier, _ = report
    done = build_report_snapshot(repo, identifier, thresholds=Thresholds(0, 999, 999))
    assert done.detection_status == NO_EVENTS and len(done.detection.events) == 0
    assert build_report_snapshot(repo, identifier).detection_status == NOT_DETECTED


def test_pdf_valid_chinese_metadata_statistics_vector_charts_and_pages(report, tmp_path):
    snapshot = report[2]
    path = tmp_path / "中文专业报告.pdf"
    export_pdf(snapshot, path)
    assert path.read_bytes().startswith(b"%PDF-")
    pdf = PdfReader(path)
    texts = [page.extract_text() for page in pdf.pages]
    full = "\n".join(texts)
    for text in (APP_NAME, snapshot.analysis.dataset_name, snapshot.report_id, "模拟数据，非真实设备测量",
                 "电压有效值时间曲线", "电流有效值时间曲线", "单相功率曲线", "算术平均值", "220", "异常事件总数"):
        assert text in full
    assert len(pdf.pages) >= 7
    assert all(f"第 {j + 1} 页" in text for j, text in enumerate(texts))
    assert pdf.metadata.title == snapshot.title
    assert any("STSong" in str(page["/Resources"].get("/Font")) or b"STSong" in path.read_bytes() for page in pdf.pages)
    # Curves are actual vector paths, not placeholder images or UI screenshots.
    assert sum(page.get_contents().get_data().count(b" l") for page in pdf.pages) >= 40


def test_pdf_long_anomaly_table_repeats_headers(report, tmp_path):
    repo = report[0]
    identifier = repo.save(generate_demo(250), "分页模拟数据")
    snapshot = build_report_snapshot(repo, identifier, thresholds=Thresholds(0, 1, 0.1))
    path = tmp_path / "多页.pdf"
    export_pdf(snapshot, path)
    pdf = PdfReader(path)
    event_pages = [page.extract_text() for page in pdf.pages if "原始索引" in page.extract_text() and "采样时间" in page.extract_text()]
    assert len(event_pages) > 2
    assert all("用户阈值" in text for text in event_pages)
    assert "249" in "\n".join(event_pages)


@pytest.mark.parametrize("kind", ["pdf", "xlsx"])
def test_export_default_no_overwrite_and_failure_preserves_old(report, tmp_path, monkeypatch, kind):
    exporter = export_pdf if kind == "pdf" else export_excel
    path = tmp_path / f"已有文件.{kind}"
    path.write_bytes(b"original")
    with pytest.raises(DataError, match="确认覆盖"):
        exporter(report[2], path)
    def fail(*args):
        raise OSError("磁盘写入失败（测试）")
    monkeypatch.setattr("smartpower.reporting.files.os.replace", fail)
    with pytest.raises(DataError, match="已有文件"):
        exporter(report[2], path, overwrite=True)
    assert path.read_bytes() == b"original"
    assert sorted(p.name for p in tmp_path.glob(f"*.{kind}")) == [path.name]


@pytest.mark.parametrize("kind", ["pdf", "xlsx"])
def test_wrong_extension_and_missing_directory(report, tmp_path, kind):
    exporter = export_pdf if kind == "pdf" else export_excel
    with pytest.raises(DataError, match="扩展名"):
        exporter(report[2], tmp_path / "报告.txt")
    with pytest.raises(DataError, match="目录不存在"):
        exporter(report[2], tmp_path / "不存在" / f"报告.{kind}")
    assert not (tmp_path / "不存在").exists()


def test_font_missing_is_explicit(report, tmp_path):
    with pytest.raises(DataError, match="字体不可用"):
        export_pdf(report[2], tmp_path / "字体.pdf", font_path=str(tmp_path / "不存在.ttf"))
    assert not (tmp_path / "字体.pdf").exists()


def test_latin_only_font_is_rejected_before_missing_glyph_pdf(report, tmp_path):
    from pathlib import Path
    import reportlab
    font = Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"
    assert font.is_file()
    with pytest.raises(DataError, match="缺少报告所需字形"):
        export_pdf(report[2], tmp_path / "缺字.pdf", font_path=str(font))
    assert not (tmp_path / "缺字.pdf").exists()


def test_pdf_reference_rejects_unsupported_nonbmp_characters(report, tmp_path):
    identifier = report[0].save(generate_demo(1), "数据😀")
    snapshot = build_report_snapshot(report[0], identifier)
    with pytest.raises(DataError, match="中文字体缺少"):
        export_pdf(snapshot, tmp_path / "非BMP.pdf")


def test_nanosecond_dates_are_preserved_as_exact_text(report, tmp_path):
    import pandas as pd
    dataset = generate_demo(2)
    dataset.frame["时间"] = pd.to_datetime(["2026-01-02 00:00:00.123456789", "2026-01-01 00:00:00.123456788"], format="mixed")
    identifier = report[0].save(dataset, "纳秒且乱序")
    snapshot = build_report_snapshot(report[0], identifier)
    path = tmp_path / "精度.xlsx"
    export_excel(snapshot, path)
    book = load_workbook(path)
    raw = list(book["原始数据"].iter_rows(values_only=True))
    assert raw[1][0].day == 2 and raw[2][0].day == 1
    assert raw[1][5].endswith("123456789") and raw[2][5].endswith("123456788")
    book.close()


def test_xlsx_out_of_range_dates_fail_without_truncation(report, tmp_path):
    import pandas as pd
    dataset = generate_demo(1)
    dataset.frame["时间"] = pd.to_datetime(["1899-12-31"])
    identifier = report[0].save(dataset, "旧日期")
    with pytest.raises(DataError, match="1900"):
        export_excel(build_report_snapshot(report[0], identifier), tmp_path / "旧日期.xlsx")


@pytest.mark.parametrize("title", ["", "a" * 121, "含\n控制符"])
def test_invalid_report_title_is_rejected(report, title):
    with pytest.raises(DataError, match="报告标题"):
        build_report_snapshot(report[0], report[1], title=title)


def test_xlsx_valid_complete_order_types_statistics_events(report, tmp_path):
    snapshot = report[2]
    path = tmp_path / "中文分析.xlsx"
    export_excel(snapshot, path)
    book = load_workbook(path, data_only=False)
    assert book.sheetnames == ["报告概况", "电压分析", "电流分析", "功率分析", "异常检测", "原始数据"]
    overview = {row[0]: row[1] for row in book["报告概况"].iter_rows(values_only=True)}
    assert overview["软件名称"] == APP_NAME and overview["数据集名称"] == snapshot.analysis.dataset_name
    raw = list(book["原始数据"].iter_rows(values_only=True))
    assert len(raw) == snapshot.analysis.count + 1
    assert [row[1] for row in raw[1:]] == list(snapshot.analysis.samples.voltage)
    assert all(row[0].year == 2026 and isinstance(row[1], (float, int)) for row in raw[1:])
    assert all("模拟数据" in row[4] for row in raw[1:])
    assert book["原始数据"].freeze_panes == "A2" and book["原始数据"].auto_filter.ref == "A1:F6"
    voltage = list(book["电压分析"].iter_rows(values_only=True))
    assert voltage[1][2] == pytest.approx(snapshot.analysis.voltage.mean)
    assert book["电压分析"]["A1"].font.bold
    anomaly = list(book["异常检测"].iter_rows(values_only=True))
    assert len(anomaly[13:]) == len(snapshot.detection.events)
    assert [r[2] for r in anomaly[13:]] == [e.row_index for e in snapshot.detection.events]
    power = list(book["功率分析"].iter_rows(values_only=True))
    header = next(j for j, row in enumerate(power) if row[0] == "时间")
    assert power[header + 1][1] == pytest.approx(snapshot.analysis.power.active[0])
    book.close()


def test_xlsx_not_executed_does_not_claim_zero(report, tmp_path):
    snapshot = build_report_snapshot(report[0], report[1])
    path = tmp_path / "未检测.xlsx"
    export_excel(snapshot, path)
    book = load_workbook(path)
    cells = [row for row in book["异常检测"].iter_rows(values_only=True)]
    assert any(NOT_DETECTED in str(row) for row in cells)
    assert not any("异常事件总数" in str(row) for row in cells)
    book.close()


def test_xlsx_row_limit_rejects_no_truncation(report, tmp_path, monkeypatch):
    monkeypatch.setattr("smartpower.reporting.excel_exporter.MAX_ROWS", 20)
    with pytest.raises(DataError, match="行数容量"):
        export_excel(report[2], tmp_path / "过量.xlsx")
    assert not (tmp_path / "过量.xlsx").exists()


@pytest.mark.parametrize("text", ["=1+1", "+cmd", "-1+2", "@SUM(A1:A3)", "<b>用户标题</b>"])
def test_untrusted_text_is_literal_xlsx_and_escaped_pdf(report, tmp_path, text):
    repo = report[0]
    identifier = repo.save(generate_demo(1), text)
    snapshot = build_report_snapshot(repo, identifier, title=text)
    path = tmp_path / "文本.xlsx"
    export_excel(snapshot, path)
    book = load_workbook(path, data_only=False)
    assert not any(cell.data_type == "f" for ws in book for row in ws for cell in row)
    assert any(cell.value == text for row in book["报告概况"] for cell in row)
    book.close()
    export_pdf(snapshot, tmp_path / "文本.pdf")
    assert text in "\n".join(p.extract_text() for p in PdfReader(tmp_path / "文本.pdf").pages)
