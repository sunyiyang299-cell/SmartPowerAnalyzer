"""Release hardening: isolated end-to-end, interrupted processes and I/O failures."""
import errno
import json
import os
from pathlib import Path
import socket
import sqlite3
import subprocess
import sys
import time

import numpy as np
from openpyxl import load_workbook
from pypdf import PdfReader
import pytest
from PySide6.QtCore import QEventLoop, QTimer

from smartpower.application import create_application
from smartpower.analysis.service import analyze_saved_dataset
from smartpower.anomaly.models import Thresholds
from smartpower.data.demo import export_demo, generate_demo
from smartpower.data.importer import import_file
from smartpower.data.validation import DataError
from smartpower.reporting.files import atomic_report
from smartpower.reporting.service import build_report_snapshot
from smartpower.reporting.pdf_exporter import export_pdf
from smartpower.reporting.excel_exporter import export_excel
from smartpower.settings.manager import SettingsManager
from smartpower.storage.repository import DatasetRepository
from smartpower.ui.main_window import MainWindow
from smartpower.ui.tasks import TaskThread
from smartpower.visualization.plot_widget import TimePlotWidget


def test_offline_end_to_end_and_fresh_process_reopen(tmp_path, monkeypatch):
    # This denies Python sockets; it is not a physical network disconnection.
    def no_network(*args, **kwargs):
        raise AssertionError('核心业务不应请求网络')
    monkeypatch.setattr(socket.socket, 'connect', no_network)
    monkeypatch.setattr(socket, 'create_connection', no_network)
    app = create_application([])
    shell = MainWindow()
    shell.show()
    app.processEvents()
    assert shell.stack.count() == 8
    shell.close()
    shell.deleteLater()
    # This host has legacy MAX_PATH behavior. Explicit extended-length paths
    # exercise real >260-character I/O without changing machine policy.
    directory = Path('\\\\?\\' + str(tmp_path.resolve())) if os.name == 'nt' else tmp_path
    while len(str(directory / '模拟输入.csv')) < 270:
        directory /= '中文目录' * 12
        directory.mkdir()
    csv = directory / '模拟输入.csv'
    export_demo(generate_demo(150), csv)
    imported = import_file(csv)
    repo = DatasetRepository(directory / '隔离.sqlite3')
    identifier = repo.save(imported, '端到端模拟样本')
    analysis = analyze_saved_dataset(repo, identifier)
    snapshot = build_report_snapshot(repo, identifier, thresholds=Thresholds(220, 240, 20))
    chart = TimePlotWidget('voltage')
    try:
        chart.resize(900, 620)
        chart.set_result(analysis)
        chart.set_detection(snapshot.detection)
        chart.show()
        app.processEvents()
        assert chart.grab().save(str(directory / '曲线.png'))
        np.testing.assert_array_equal(next(iter(chart.curves.values())).getData()[1], imported.frame['电压'].to_numpy())
    finally:
        chart.close()
        chart.deleteLater()
    pdf, excel = directory / '报告.pdf', directory / '报告.xlsx'
    export_pdf(snapshot, pdf)
    export_excel(snapshot, excel)
    text = '\n'.join(page.extract_text() for page in PdfReader(pdf).pages)
    assert '端到端模拟样本' in text and '模拟数据，非真实设备测量' in text
    assert str(len(snapshot.detection.events)) in text
    book = load_workbook(excel, read_only=True, data_only=False)
    try:
        assert len(book.sheetnames) == 6
        rows = list(book['原始数据'].values)
        # All source records must appear, independently of plot reduction.
        from datetime import datetime
        assert len(rows) == 151
        assert all(isinstance(row[0], datetime) for row in rows[1:])
        np.testing.assert_allclose([row[1] for row in rows[1:]], imported.frame['电压'], rtol=1e-14, atol=0)
        assert '模拟数据' in str(rows)
    finally:
        book.close()
    code = '''import json, sys
from smartpower.storage.repository import DatasetRepository
r=DatasetRepository(sys.argv[1]); d=r.load(int(sys.argv[2]))
print(json.dumps([len(d.frame),d.simulated,d.name]))
'''
    child = subprocess.run([sys.executable, '-c', code, str(repo.path), str(identifier)], capture_output=True, text=True, encoding='utf-8', timeout=30)
    assert child.returncode == 0, child.stderr
    assert json.loads(child.stdout) == [150, True, '端到端模拟样本']
    assert repo.load(identifier).frame.equals(imported.frame)


@pytest.mark.parametrize('operation', ['import', 'database', 'analysis', 'pdf', 'excel'])
def test_close_protection_while_real_background_task_is_pending(operation):
    app = create_application([])
    window = MainWindow()
    key = {'import':'data', 'database':'data', 'analysis':'voltage', 'pdf':'reports', 'excel':'reports'}[operation]
    page = window.pages[key]
    # Hold an actual QThread open without writing any user files.
    import threading
    released = threading.Event()
    thread = TaskThread(lambda: released.wait(10), page)
    page._thread = thread
    window.show()
    thread.start()
    try:
        assert window.close() is False
        assert window.isVisible() and window.stack.currentWidget() is page
        assert '后台任务' in (page.message.toPlainText() if hasattr(page.message, 'toPlainText') else page.message.text())
    finally:
        released.set()
        assert thread.wait(10000)
        page._thread = None
        window.close()
        window.deleteLater()
        app.processEvents()


@pytest.mark.parametrize('kind', ['database', 'report'])
def test_killed_writer_preserves_original_and_restart_integrity(tmp_path, kind):
    repo = DatasetRepository(tmp_path / '重启.sqlite3')
    identifier = repo.save(generate_demo(10), '原始模拟样本')
    target = tmp_path / '已有.pdf'
    target.write_bytes(b'original target')
    marker = tmp_path / 'ready'
    code = '''import sqlite3, sys, time
from pathlib import Path
from smartpower.reporting.files import atomic_report
kind, database, output, ready = sys.argv[1:]
if kind == 'database':
    c=sqlite3.connect(database); c.execute('BEGIN IMMEDIATE')
    c.execute("UPDATE datasets SET name='uncommitted'")
    Path(ready).write_text('ready')
    time.sleep(60)
else:
    with atomic_report(output, '.pdf', overwrite=True) as temporary:
        temporary.write_bytes(b'partial uncommitted report')
        Path(ready).write_text(str(temporary), encoding='utf-8')
        time.sleep(60)
'''
    child = subprocess.Popen([sys.executable, '-c', code, kind, str(repo.path), str(target), str(marker)], stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    try:
        deadline = time.monotonic() + 20
        while not marker.exists() and time.monotonic() < deadline and child.poll() is None:
            time.sleep(.02)
        assert marker.exists(), child.communicate(timeout=2) if child.poll() is not None else '子进程未准备好'
    finally:
        child.kill()
        child.communicate(timeout=10)
    assert repo.load(identifier).name == '原始模拟样本'
    with sqlite3.connect(repo.path) as connection:
        assert connection.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
    assert target.read_bytes() == b'original target'
    if kind == 'report':
        # OS termination cannot run finally. Clean only this test's known temp.
        leftover = Path(marker.read_text(encoding='utf-8'))
        assert leftover.parent == tmp_path and leftover != target
        assert leftover.read_bytes() == b'partial uncommitted report'
        leftover.unlink()


def test_simulated_disk_full_preserves_settings_and_report(tmp_path, monkeypatch):
    manager = SettingsManager(tmp_path / 'settings.json', tmp_path)
    manager.save(manager.defaults)
    original = manager.path.read_bytes()
    def disk_full(*args, **kwargs):
        raise OSError(errno.ENOSPC, '测试模拟磁盘空间不足')
    monkeypatch.setattr('smartpower.settings.manager.os.replace', disk_full)
    with pytest.raises(DataError, match='旧设置文件保留'):
        manager.save(manager.defaults)
    assert manager.path.read_bytes() == original
    target = tmp_path / 'existing.pdf'
    target.write_bytes(b'old')
    with pytest.raises(DataError, match='空间不足'):
        with atomic_report(target, '.pdf', overwrite=True) as temporary:
            temporary.write_bytes(b'new')
    assert target.read_bytes() == b'old'
    assert sorted(p.name for p in tmp_path.iterdir()) == ['existing.pdf', 'settings.json']


@pytest.mark.skipif(os.name != 'nt', reason='仅适用 Windows 文件共享模式')
def test_actual_windows_locked_target_and_readonly_target(tmp_path):
    import ctypes
    from ctypes import wintypes
    target = tmp_path / '占用.pdf'
    target.write_bytes(b'original')
    kernel = ctypes.WinDLL('kernel32', use_last_error=True)
    kernel.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, ctypes.c_void_p, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
    kernel.CreateFileW.restype = wintypes.HANDLE
    kernel.CloseHandle.argtypes = [wintypes.HANDLE]
    handle = kernel.CreateFileW(str(target), 0x80000000, 0, None, 3, 0x80, None)
    assert handle != wintypes.HANDLE(-1).value
    try:
        with pytest.raises(DataError, match='未被本次失败写入替换'):
            with atomic_report(target, '.pdf', overwrite=True) as temporary:
                temporary.write_bytes(b'new')
    finally:
        assert kernel.CloseHandle(handle)
    assert target.read_bytes() == b'original'
    import stat
    target.chmod(stat.S_IREAD)
    try:
        with pytest.raises(DataError):
            with atomic_report(target, '.pdf', overwrite=True) as temporary:
                temporary.write_bytes(b'new')
        assert target.read_bytes() == b'original'
    finally:
        target.chmod(stat.S_IWRITE | stat.S_IREAD)
