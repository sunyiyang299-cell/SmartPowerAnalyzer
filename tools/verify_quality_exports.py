"""Verify complete benchmark artifacts; never open the user's database."""
import argparse
from datetime import datetime
import json
import math
from pathlib import Path
import re
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from pypdf import PdfReader
from smartpower.anomaly.detector import detect_anomalies
from smartpower.anomaly.models import Thresholds
from smartpower.storage.repository import DatasetRepository


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('directory', type=Path)
    args = parser.parse_args()
    root = args.directory
    dataset = DatasetRepository(root / 'isolated.sqlite3').load(1)
    assert dataset.simulated
    detection = detect_anomalies(dataset, Thresholds(220, 240, 20), 1)
    reader = PdfReader(root / '报告.pdf')
    text = '\n'.join(p.extract_text() for p in reader.pages)
    actual = re.findall(r'(\d+)\n(\d{4}-\d{2}-\d{2} [\d:.]+)\n(电压偏低|电压偏高|电流偏高)\n([\d.e+\-]+)\n([\d.e+\-]+)\n([VA])\n', text)
    expected = [(str(e.row_index), str(e.timestamp).replace('T', ' '), e.kind,
                 f'{e.value:.12g}', f'{e.threshold:.12g}', e.unit) for e in detection.events]
    assert actual == expected, (len(actual), len(expected))
    for title in ['电压有效值时间曲线', '电流有效值时间曲线', '单相功率曲线', '模拟数据，非真实设备测量']:
        assert title in text
    for i, page in enumerate(reader.pages, 1):
        assert f'第 {i} 页' in page.extract_text()
    book = load_workbook(root / '报告.xlsx', read_only=True, data_only=False)
    try:
        assert book.sheetnames == ['报告概况', '电压分析', '电流分析', '功率分析', '异常检测', '原始数据']
        count = 0
        for pos, row in enumerate(book['原始数据'].iter_rows(min_row=2, values_only=True)):
            assert isinstance(row[0], datetime)
            time, voltage, current, pf = dataset.frame.iloc[pos]
            assert row[0] == time.to_pydatetime()
            assert all(math.isclose(value, reference, rel_tol=1e-14, abs_tol=0) for value, reference in zip(row[1:4], (voltage, current, pf)))
            assert row[4] == '模拟数据，非真实设备测量'
            count += 1
        assert count == len(dataset.frame)
        event_count = 0
        for row, e in zip(book['异常检测'].iter_rows(min_row=14, values_only=True), detection.events, strict=True):
            assert row[0] == 1 and row[2] == e.row_index and row[5] == e.kind
            assert row[7] == e.unit and row[8] == e.source
            assert math.isclose(row[4], e.value, rel_tol=1e-14, abs_tol=0)
            event_count += 1
        assert event_count == len(detection.events)
    finally:
        book.close()
    print(json.dumps({'records':count, 'pdf_pages':len(reader.pages), 'verified_pdf_events':len(actual), 'verified_excel_events':event_count, 'result':'PASS'}))


if __name__ == '__main__':
    main()
