"""Styled streaming XLSX with full original-order samples and literal text."""

from datetime import datetime
from math import ceil
from unicodedata import east_asian_width

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from smartpower.analysis.power import POWER_MODEL
from smartpower.data.validation import DataError
from smartpower.reporting.files import atomic_report
from smartpower.reporting.models import LIMITATIONS, statistics_rows

MAX_ROWS = 1_048_576


def excel_time(value):
    stamp = pd.Timestamp(value)
    if stamp.year < 1900 or stamp.year > 9999:
        raise DataError("Excel 原生日期支持本报告的 1900～9999 年范围，当前时间超出范围，已拒绝导出，未截断数据。")
    return stamp.to_pydatetime(warn=False)


def export_excel(snapshot, path, *, overwrite=False):
    a = snapshot.analysis
    events = snapshot.detection.events if snapshot.detection is not None else ()
    # Power sheet: 13 statistics + heading/source/sample-header = 16 rows.
    # Detection sheet: 13 rows before its first event; overview has 23 rows.
    if max(a.count + 16, len(events) + 13, 23) > MAX_ROWS:
        raise DataError("数据超过 Excel 单工作表行数容量（含表头和统计区），已拒绝导出；未截断原始数据或异常事件。")
    # Fail before publication if any native date cannot be represented.
    excel_time(snapshot.time_start)
    excel_time(snapshot.time_end)
    with atomic_report(path, ".xlsx", overwrite=overwrite) as temporary:
        _write_workbook(snapshot, temporary)


def _write_workbook(snapshot, temporary):
    a = snapshot.analysis
    events = snapshot.detection.events if snapshot.detection is not None else ()
    book = Workbook(write_only=True)
    row_numbers = {}
    try:

        def sheet(name, widths):
            ws = book.create_sheet(name)
            for j, width in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(j)].width = width
            return ws

        def append(ws, values, *, header=False):
            cells = []
            lines = 1
            for column, value in enumerate(values, 1):
                if isinstance(value, np.generic):
                    value = value.item()
                cell = WriteOnlyCell(ws, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"  # Literal text even when starting with =,+,-,@.
                    width = ws.column_dimensions[get_column_letter(column)].width - 2
                    text_width = sum(2 if east_asian_width(char) in "WF" else 1 for char in value)
                    lines = max(lines, ceil(text_width / max(width, 1)))
                elif isinstance(value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm:ss.000"
                elif isinstance(value, (float, int)):
                    cell.number_format = "0.###############"
                if header:
                    cell.font = Font(bold=True, color="183153")
                    cell.fill = PatternFill("solid", fgColor="E9F0FA")
                cell.alignment = Alignment(vertical="top", wrap_text=isinstance(value, str))
                cells.append(cell)
            row_numbers[ws.title] = row_numbers.get(ws.title, 0) + 1
            ws.row_dimensions[row_numbers[ws.title]].height = max(20, lines * 15)
            ws.append(cells)

        overview = sheet("报告概况", (28, 105))
        overview.freeze_panes = "A2"
        append(overview, ("项目", "内容"), header=True)
        for row in (("报告唯一标识", snapshot.report_id), ("软件名称", snapshot.product_name), ("开发版本", snapshot.version),
                    ("报告标题", snapshot.title), ("数据集 ID", a.dataset_id), ("数据集名称", a.dataset_name),
                    ("数据来源", snapshot.source_label), ("采样数量", a.count), ("起始时间", excel_time(snapshot.time_start)),
                    ("结束时间", excel_time(snapshot.time_end)), ("生成时间", snapshot.created_at.isoformat()),
                    ("计算模型", POWER_MODEL), ("检测状态", snapshot.detection_status),
                    ("日期精度", "日期列为 Excel 原生日期；Excel 保存精度有限，原始纳秒精度时间另以精确时间文本列保留。")):
            append(overview, row)
        for text in LIMITATIONS:
            append(overview, ("说明与限制", text))
        stats = statistics_rows(snapshot)
        for name, quantities, values, headers in (
            ("电压分析", {a.voltage.quantity}, (a.samples.voltage,), ("时间", "电压有效值（V）", "精确时间文本")),
            ("电流分析", {a.current.quantity}, (a.samples.current,), ("时间", "电流有效值（A）", "精确时间文本")),
            ("功率分析", {a.power.active_statistics.quantity, a.power.apparent_statistics.quantity,
                           a.power.reactive_abs_statistics.quantity, a.power.power_factor_statistics.quantity},
             (a.power.active, a.power.apparent, a.power.reactive_abs, snapshot.power_factors),
             ("时间", "有功功率 P（W）", "视在功率 S（VA）", "无功绝对值 |Q|（var）", "功率因数 PF（无量纲）", "精确时间文本")),
        ):
            rows = [row for row in stats if row[0] in quantities]
            header_row = len(rows) + 3
            ws = sheet(name, (28, 30, 30, 30, 30, 36))
            ws.freeze_panes = f"A{header_row + 1}"
            append(ws, ("参数", "统计指标", "数值", "单位"), header=True)
            for row in rows:
                append(ws, row)
            append(ws, (snapshot.source_label,))
            append(ws, headers, header=True)
            for pos in range(a.count):
                append(ws, (excel_time(a.times[pos]), *(float(v[pos]) for v in values), str(a.times[pos]).replace("T", " ")))
            ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{header_row + a.count}"
        anomaly = sheet("异常检测", (14, 28, 22, 14, 22, 18, 22, 10, 40, 36))
        anomaly.freeze_panes = "A2"
        append(anomaly, ("检测项目", "内容"), header=True)
        append(anomaly, ("检测状态", snapshot.detection_status))
        append(anomaly, ("来源", snapshot.source_label))
        if snapshot.detection is not None:
            d, t = snapshot.detection, snapshot.thresholds
            for row in (("电压下限（V）", t.voltage_lower), ("电压上限（V）", t.voltage_upper),
                        ("电流上限（A）", t.current_upper), ("异常事件总数", len(d.events)),
                        ("唯一异常采样点", d.unique_sample_count), *d.type_counts.items()):
                append(anomaly, row)
            append(anomaly, ("说明", "仅按用户阈值检测，等号不触发；阈值不是国家标准限值。"))
            header_row = 13
            append(anomaly, ("数据集 ID", "时间", "原始索引（从0起）", "参数", "采样值", "异常类型", "阈值", "单位", "来源", "精确时间文本"), header=True)
            for e in events:
                append(anomaly, (e.dataset_id, excel_time(e.timestamp), e.row_index, e.parameter, e.value,
                                  e.kind, e.threshold, e.unit, e.source, str(e.timestamp).replace("T", " ")))
            anomaly.freeze_panes = f"A{header_row + 1}"
            anomaly.auto_filter.ref = f"A{header_row}:J{header_row + len(events)}"
        raw = sheet("原始数据", (28, 24, 24, 26, 40, 36))
        raw.freeze_panes = "A2"
        append(raw, ("时间", "电压（V）", "电流（A）", "功率因数（无量纲）", "数据来源", "精确时间文本"), header=True)
        for pos in range(a.count):
            append(raw, (excel_time(a.times[pos]), float(a.samples.voltage[pos]), float(a.samples.current[pos]),
                         float(snapshot.power_factors[pos]), snapshot.source_label, str(a.times[pos]).replace("T", " ")))
        raw.auto_filter.ref = f"A1:F{a.count + 1}"
        book.save(temporary)
    finally:
        for ws in book.worksheets:
            if not ws.closed:
                ws.close()
        book.close()
