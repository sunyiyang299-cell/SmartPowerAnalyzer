"""Explicit isolated release verification for source and frozen Windows builds.

Only writes to a newly created verification directory. No arbitrary scripts are
loaded; re-open only accepts directories bearing this verifier's marker.
"""
from dataclasses import replace
import json
import os
from pathlib import Path
import socket
import sys
import traceback

OWNER = 'SmartPowerAnalyzer-release-verification-v1'


def run_release_check(directory, *, reopen=False):
    root = Path(directory).resolve()
    if reopen:
        metadata = json.loads((root / 'verification.json').read_text(encoding='utf-8'))
        if metadata.get('owner') != OWNER:
            raise ValueError('只能重新打开由发行验证创建的隔离目录。')
    else:
        root.mkdir(parents=True, exist_ok=False)
        (root / 'verification.json').write_text(json.dumps({'owner':OWNER}), encoding='utf-8')
    os.environ['SMARTPOWER_DB_PATH'] = str(root / 'isolated.sqlite3')
    os.environ['SMARTPOWER_SETTINGS_PATH'] = str(root / 'settings.json')
    # Deny Python outbound sockets. This is not an OS-level physical disconnect.
    def deny_network(*args, **kwargs):
        raise RuntimeError('发行验证禁止网络连接')
    socket.socket.connect = deny_network
    socket.create_connection = deny_network
    from PySide6.QtCore import QEventLoop, QTimer, Qt
    from PySide6.QtWidgets import QFileDialog, QLabel, QLineEdit, QPlainTextEdit
    from smartpower.application import create_application
    from smartpower.config import APP_NAME, APP_VERSION
    from smartpower.data.demo import generate_demo, export_demo
    from smartpower.data.validation import SOURCE_COLUMN, SIMULATED_SOURCE
    from smartpower.storage.repository import DatasetRepository
    from smartpower.ui.main_window import MainWindow
    from smartpower.ui.paths import default_settings_manager
    from smartpower.visualization.export import save_png
    import numpy as np
    from openpyxl import load_workbook

    app = create_application([])
    app.setAttribute(Qt.ApplicationAttribute.AA_DontUseNativeDialogs, True)
    if not reopen:
        manager = default_settings_manager()
        manager.save(replace(manager.defaults, report_directory=str(root), chart_directory=str(root)))
    window = MainWindow()
    window.show()
    result = {'version':APP_VERSION, 'product':APP_NAME, 'frozen':bool(getattr(sys,'frozen',False)),
              'platform':app.platformName(), 'socket_network':'disabled', 'steps':[]}
    result_path = root / ('reopen-result.json' if reopen else 'result.json')

    def step(name):
        result['steps'].append(name)
        result_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')

    def wait_idle(page):
        if page.busy:
            loop = QEventLoop()
            timeout = QTimer()
            timeout.setSingleShot(True)
            timeout.timeout.connect(loop.quit)
            page.operation_finished.connect(loop.quit)
            timeout.start(120000)
            loop.exec()
            timeout.stop()
            page.operation_finished.disconnect(loop.quit)
        assert not page.busy, '后台任务超时'

    def screenshot(key):
        window.navigate(key)
        page = window.pages[key]
        page.verticalScrollBar().setValue(0)
        edits = []
        # Redact paths on actual widgets, preserving real computed data and UI.
        for widget in page.findChildren(QLabel) + page.findChildren(QPlainTextEdit):
            plain = isinstance(widget, QPlainTextEdit)
            text = widget.toPlainText() if plain else widget.text()
            if str(root) in text:
                edits.append((widget, text, plain))
                replaced = text.replace(str(root), '（隔离验证目录）')
                widget.setPlainText(replaced) if plain else widget.setText(replaced)
        fields = [edit for edit in page.findChildren(QLineEdit) if str(root) in edit.text()]
        for field in fields:
            field.setEchoMode(QLineEdit.EchoMode.Password)
        app.processEvents()
        shots = root / 'screenshots'
        shots.mkdir(exist_ok=True)
        assert window.grab().save(str(shots / f'{key}.png'))
        for widget, text, plain in edits:
            widget.setPlainText(text) if plain else widget.setText(text)
        for field in fields:
            field.setEchoMode(QLineEdit.EchoMode.Normal)

    def choose_report(kind):
        timer = QTimer()
        timer.setInterval(50)
        def accept():
            for dialog in app.topLevelWidgets():
                if isinstance(dialog, QFileDialog) and dialog.isVisible():
                    dialog.setDirectory(str(root))
                    dialog.findChild(QLineEdit, 'fileNameEdit').setText(f'发行验证报告.{kind}')
                    timer.stop()
                    dialog.accept()
                    return
        timer.timeout.connect(accept)
        timer.start()
        window.pages['reports'].start_export(kind)
        timer.stop()
        wait_idle(window.pages['reports'])
        assert window.pages['reports'].current_snapshot is not None
        assert (root / f'发行验证报告.{kind}').stat().st_size > 1000

    def exercise():
        try:
            assert window.windowTitle() == APP_NAME and window.isVisible()
            assert window.stack.count() == 8
            for key in window.pages:
                window.navigation_buttons[key].click()
                app.processEvents()
                assert window.stack.currentWidget() is window.pages[key]
            step('eight_pages_native_navigation')
            repository = DatasetRepository(root / 'isolated.sqlite3')
            if reopen:
                saved = repository.list_datasets()
                assert len(saved) == 2
                assert all(item['row_count'] == 1000 and item['simulated'] for item in saved)
                assert default_settings_manager().load().value.startup_page == 'reports'
                assert window.stack.currentWidget() is window.pages['settings']  # Last navigation above.
                assert len(repository.load(saved[0]['id']).frame) == 1000
                step('database_and_settings_preserved_after_restart')
            else:
                demo = generate_demo()
                csv = root / '模拟输入.csv'
                excel = root / '模拟输入.xlsx'
                export_demo(demo, csv)
                frame = demo.frame.copy()
                frame[SOURCE_COLUMN] = SIMULATED_SOURCE
                frame.to_excel(excel, index=False)
                data = window.pages['data']
                for path, name in [(csv,'发行验证模拟数据'), (excel,'Excel模拟导入')]:
                    data.import_path(path)
                    wait_idle(data)
                    assert data.dataset.simulated and data.model.rowCount() == 1000
                    data.name_edit.setText(name)
                    data.save()
                    wait_idle(data)
                    assert not data.has_unsaved_data
                step('csv_excel_import_and_sqlite_save')
                data.refresh_history(); wait_idle(data)
                data.history_combo.setCurrentIndex(data.history_combo.findData(1))
                data.load_selected(); wait_idle(data)
                assert data.dataset.name == '发行验证模拟数据'
                step('sqlite_history_read')
                for key in ['voltage','current','power']:
                    page = window.pages[key]
                    page.refresh_datasets(); wait_idle(page)
                    page.dataset_combo.setCurrentIndex(page.dataset_combo.findData(1))
                    page.analyze_button.click(); wait_idle(page)
                    assert page.current_result.count == 1000 and page.current_result.simulated
                    assert page.chart.curves
                    if key == 'power':
                        expected = data.dataset.frame['电压'] * data.dataset.frame['电流'] * data.dataset.frame['功率因数']
                        np.testing.assert_allclose(page.current_result.power.active, expected, rtol=1e-12)
                    save_png(page.chart.grab(), root / f'{key}.png')
                step('voltage_current_power_analysis_curves_and_png')
                anomaly = window.pages['anomalies']
                anomaly.refresh_datasets(); wait_idle(anomaly)
                anomaly.dataset_combo.setCurrentIndex(anomaly.dataset_combo.findData(1))
                anomaly.apply_demo_preset(); anomaly.enabled_check.setChecked(True)
                anomaly.analyze_button.click(); wait_idle(anomaly)
                values = data.dataset.frame
                expected_events = int((values['电压']<210).sum()+(values['电压']>230).sum()+(values['电流']>20).sum())
                assert len(anomaly.current_result.events) == expected_events
                result['events'] = expected_events
                step('full_anomaly_detection_and_plot_link')
                reports = window.pages['reports']
                reports.refresh_datasets(); wait_idle(reports)
                reports.dataset_combo.setCurrentIndex(reports.dataset_combo.findData(1))
                reports.detection_enabled.setChecked(True)
                for edit, value in zip(reports.threshold_inputs, ['210','230','20']): edit.setText(value)
                choose_report('pdf'); choose_report('xlsx')
                assert (root / '发行验证报告.pdf').read_bytes().startswith(b'%PDF-')
                book = load_workbook(root / '发行验证报告.xlsx', read_only=True)
                try:
                    assert len(book.sheetnames) == 6
                    assert sum(1 for _ in book['原始数据'].iter_rows(min_row=2)) == 1000
                finally: book.close()
                step('pdf_and_excel_export_via_real_save_dialog')
                settings = window.pages['settings']
                settings.startup_page.setCurrentIndex(settings.startup_page.findData('reports'))
                settings.save_button.click()
                assert default_settings_manager().load().value.startup_page == 'reports'
                step('settings_saved')
                for key in window.pages: screenshot(key)
                step('actual_gui_screenshots_paths_redacted')
            result['status'] = 'PASS'
            result['records'] = 1000
            result_path.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
            window.close()
            app.exit(0)
        except Exception:
            result['status'] = 'FAIL'
            result['error'] = traceback.format_exc().replace(str(root), '<verification-directory>')
            result_path.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
            # Gracefully wait for started workers; never terminate a QThread.
            for page in window.pages.values():
                if getattr(page,'busy',False): wait_idle(page)
            window.pages['data']._dirty = False
            window.close()
            app.exit(1)
    QTimer.singleShot(150, exercise)
    return app.exec()
